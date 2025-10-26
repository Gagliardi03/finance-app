import streamlit as st
from backend import ExpenseService
from typing import List, Dict
from datetime import datetime
from io import BytesIO
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side


def show_expenses_page(expense_service: ExpenseService):
    """
    Display expenses page for adding expenses.
    New flow: Add multiple expenses, then save all at once.
    View analytics in a separate view.
    DARK MODE PERMANENT.

    Args:
        expense_service: Service for expense operations
    """
    # Initialize session state for temporary expenses
    if "temp_expenses" not in st.session_state:
        st.session_state.temp_expenses = []
    if "show_analytics" not in st.session_state:
        st.session_state.show_analytics = False

    # Page header - DARK
    st.markdown(
        """
        <div style="text-align: center; padding: 32px 24px; background: #1a1a1a; 
                    border-radius: 12px; box-shadow: 0 6.4px 14.4px rgba(0, 0, 0, 0.5); 
                    margin-bottom: 32px; border: 1px solid #333333;">
            <h1 style="font-size: 32px; font-weight: 700; color: #ffffff; margin-bottom: 8px;">
                💸 Expense Tracker
            </h1>
            <p style="font-size: 16px; color: #e0e0e0; margin: 0;">
                Add your expenses and save them when ready
            </p>
        </div>
        """,
        unsafe_allow_html=True,
    )

    # Navigation buttons
    col1, col2, col3 = st.columns([1, 1, 4])
    with col1:
        if st.button("⬅️ Back to Home", use_container_width=True):
            st.session_state.page = "home"
            st.session_state.show_analytics = False
            st.rerun()
    with col2:
        analytics_label = (
            "📝 Add Expenses"
            if st.session_state.show_analytics
            else "📊 View Analytics"
        )
        if st.button(analytics_label, use_container_width=True):
            st.session_state.show_analytics = not st.session_state.show_analytics
            st.rerun()

    st.divider()

    # ==================== ANALYTICS VIEW ====================
    if st.session_state.show_analytics:
        _show_analytics_view(expense_service)
        return

    # ==================== ADD EXPENSES VIEW ====================
    _show_add_expenses_view(expense_service)


def _show_add_expenses_view(expense_service: ExpenseService):
    """
    Show the view for adding expenses.
    DARK MODE PERMANENT.

    Args:
        expense_service: Service for expense operations
    """
    # ==================== ADD EXPENSE FORM ====================
    st.markdown(
        """
        <h2 style="font-size: 20px; font-weight: 600; color: #ffffff; margin-bottom: 16px;">
            ➕ Add New Expense
        </h2>
        """,
        unsafe_allow_html=True,
    )

    with st.form("expense_form", clear_on_submit=True):
        col1, col2, col3 = st.columns(3)

        with col1:
            # Category input
            category = st.text_input(
                "Category",
                placeholder="e.g., Pharmacy, Food, Transport",
                help="What type of expense?",
            )

        with col2:
            # Location input
            location = st.text_input(
                "Location",
                placeholder="e.g., DROGARIA PAULISTA",
                help="Where did you spend?",
            )

        with col3:
            # Price input
            price = st.number_input(
                "Price (BRL)",
                min_value=0.01,
                step=0.01,
                format="%.2f",
                help="Amount spent",
            )

        col4, col5, col6 = st.columns(3)

        with col4:
            # Payment method selection
            payment_method = st.selectbox(
                "Payment Method",
                options=["Credit Card", "Debit Card", "PIX", "Cash", "Bank Transfer"],
                help="How did you pay?",
            )

        with col5:
            # Installments (only for credit card)
            if payment_method == "Credit Card":
                installments = st.number_input(
                    "Installments",
                    min_value=1,
                    max_value=24,
                    value=1,
                    step=1,
                    help="Number of installments",
                )
            else:
                installments = 1
                st.markdown(
                    """
                    <div style="margin-top: 28px; padding: 8px 12px; background: #1b2d3d; 
                                border-radius: 4px; border-left: 3px solid #4da6ff;">
                        <p style="color: #4da6ff; font-size: 12px; margin: 0;">
                            💡 Installments only for Credit Card
                        </p>
                    </div>
                    """,
                    unsafe_allow_html=True,
                )

        with col6:
            st.write("")  # Spacing

        # Submit button
        submitted = st.form_submit_button(
            "➕ Add to List", use_container_width=True, type="primary"
        )

        if submitted:
            # Validate inputs
            if not category:
                st.error("❌ Please fill in category field!")
            else:
                # Add to temporary list
                temp_expense = {
                    "category": category,
                    "location": location if location else "-",
                    "price": price,
                    "payment_method": payment_method,
                    "installments": installments,
                    "date": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
                }
                st.session_state.temp_expenses.append(temp_expense)
                st.success(f"✅ Added: {category} - R$ {price:,.2f}")
                st.rerun()

    st.divider()

    # ==================== TEMPORARY EXPENSES LIST ====================
    if st.session_state.temp_expenses:
        st.markdown(
            f"""
            <h2 style="font-size: 20px; font-weight: 600; color: #ffffff; margin-bottom: 16px;">
                📋 Expenses to Save ({len(st.session_state.temp_expenses)})
            </h2>
            """,
            unsafe_allow_html=True,
        )

        # Display expenses in table format
        for idx, expense in enumerate(st.session_state.temp_expenses):
            col1, col2, col3, col4, col5, col6 = st.columns([2, 2, 2, 2, 2, 1])

            with col1:
                st.markdown(
                    f"""
                    <div style="padding: 8px; background: #2d2d2d; border-radius: 4px; border: 1px solid #404040;">
                        <p style="color: #e0e0e0; font-size: 10px; margin: 0; 
                                  text-transform: uppercase; letter-spacing: 0.5px;">
                            Category
                        </p>
                        <p style="color: #ffffff; font-size: 14px; font-weight: 600; margin: 4px 0 0 0;">
                            {expense['category']}
                        </p>
                    </div>
                    """,
                    unsafe_allow_html=True,
                )

            with col2:
                st.markdown(
                    f"""
                    <div style="padding: 8px; background: #2d2d2d; border-radius: 4px; border: 1px solid #404040;">
                        <p style="color: #e0e0e0; font-size: 10px; margin: 0; 
                                  text-transform: uppercase; letter-spacing: 0.5px;">
                            Location
                        </p>
                        <p style="color: #ffffff; font-size: 14px; font-weight: 600; margin: 4px 0 0 0;">
                            {expense['location']}
                        </p>
                    </div>
                    """,
                    unsafe_allow_html=True,
                )

            with col3:
                st.markdown(
                    f"""
                    <div style="padding: 8px; background: #2d2d2d; border-radius: 4px; border: 1px solid #404040;">
                        <p style="color: #e0e0e0; font-size: 10px; margin: 0; 
                                  text-transform: uppercase; letter-spacing: 0.5px;">
                            Price
                        </p>
                        <p style="color: #f44336; font-size: 14px; font-weight: 700; margin: 4px 0 0 0;">
                            R$ {expense['price']:,.2f}
                        </p>
                    </div>
                    """,
                    unsafe_allow_html=True,
                )

            with col4:
                st.markdown(
                    f"""
                    <div style="padding: 8px; background: #2d2d2d; border-radius: 4px; border: 1px solid #404040;">
                        <p style="color: #e0e0e0; font-size: 10px; margin: 0; 
                                  text-transform: uppercase; letter-spacing: 0.5px;">
                            Payment
                        </p>
                        <p style="color: #ffffff; font-size: 14px; font-weight: 600; margin: 4px 0 0 0;">
                            {expense['payment_method']}
                        </p>
                    </div>
                    """,
                    unsafe_allow_html=True,
                )

            with col5:
                st.markdown(
                    f"""
                    <div style="padding: 8px; background: #2d2d2d; border-radius: 4px; border: 1px solid #404040;">
                        <p style="color: #e0e0e0; font-size: 10px; margin: 0; 
                                  text-transform: uppercase; letter-spacing: 0.5px;">
                            Installments
                        </p>
                        <p style="color: #ffffff; font-size: 14px; font-weight: 600; margin: 4px 0 0 0;">
                            {expense['installments']}x
                        </p>
                    </div>
                    """,
                    unsafe_allow_html=True,
                )

            st.markdown("<br>", unsafe_allow_html=True)

        # Action buttons
        col1, col2, col3 = st.columns(3)

        with col1:
            # Calculate total
            total_amount = sum(e["price"] for e in st.session_state.temp_expenses)
            st.markdown(
                f"""
                <div style="padding: 12px; background: #3d1b1b; border-radius: 4px; 
                            border-left: 3px solid #f44336; text-align: center; border: 1px solid #f44336;">
                    <p style="color: #e0e0e0; font-size: 12px; margin: 0;">
                        Total Amount
                    </p>
                    <p style="color: #f44336; font-size: 20px; font-weight: 700; margin: 4px 0 0 0;">
                        R$ {total_amount:,.2f}
                    </p>
                </div>
                """,
                unsafe_allow_html=True,
            )

        with col2:
            if st.button(
                "💾 Save All Expenses",
                use_container_width=True,
                type="primary",
                key="save_all_btn",
            ):
                try:
                    # Save all expenses to database
                    for expense in st.session_state.temp_expenses:
                        expense_service.create_expense(
                            category=expense["category"],
                            location=expense["location"],
                            price=expense["price"],
                            payment_method=expense["payment_method"],
                            installments=expense["installments"],
                        )

                    # Clear temporary list
                    count = len(st.session_state.temp_expenses)
                    st.session_state.temp_expenses = []

                    st.success(f"✅ Successfully saved {count} expense(s)!")
                    st.balloons()
                    st.rerun()

                except Exception as e:
                    st.error(f"❌ Error saving expenses: {e}")

        with col3:
            if st.button(
                "🗑️ Clear All",
                use_container_width=True,
                key="clear_all_btn",
            ):
                st.session_state.temp_expenses = []
                st.success("✅ List cleared!")
                st.rerun()

    else:
        st.markdown(
            """
            <div style="padding: 48px; background: #1a1a1a; border-radius: 8px; 
                        text-align: center; border: 2px dashed #404040;">
                <div style="font-size: 48px; margin-bottom: 16px;">📭</div>
                <h3 style="color: #e0e0e0; font-size: 18px; font-weight: 600; margin: 0;">
                    No expenses added yet
                </h3>
                <p style="color: #808080; font-size: 14px; margin: 8px 0 0 0;">
                    Use the form above to add your first expense
                </p>
            </div>
            """,
            unsafe_allow_html=True,
        )


def _create_formatted_expenses_xlsx(df: pd.DataFrame) -> BytesIO:
    """
    Create a formatted Excel file for expenses data.

    Args:
        df: DataFrame with expenses data

    Returns:
        BytesIO buffer with formatted Excel file
    """
    try:
        # Create workbook and worksheet
        wb = Workbook()
        ws = wb.active
        ws.title = "Expenses"

        # Define styles
        header_font = Font(bold=True, color="FFFFFF", size=12)
        header_fill = PatternFill(
            start_color="C00000", end_color="C00000", fill_type="solid"
        )
        header_alignment = Alignment(horizontal="center", vertical="center")

        cell_alignment = Alignment(horizontal="left", vertical="center")
        number_alignment = Alignment(horizontal="right", vertical="center")

        border_style = Border(
            left=Side(style="thin"),
            right=Side(style="thin"),
            top=Side(style="thin"),
            bottom=Side(style="thin"),
        )

        # Write title
        ws.merge_cells("A1:G1")
        title_cell = ws["A1"]
        title_cell.value = "Expenses Report"
        title_cell.font = Font(bold=True, size=16, color="7F0000")
        title_cell.alignment = Alignment(horizontal="center", vertical="center")

        # Write headers
        headers = [
            "ID",
            "Date",
            "Category",
            "Location",
            "Price (R$)",
            "Payment Method",
            "Installments",
        ]
        for col_idx, header in enumerate(headers, 1):
            cell = ws.cell(row=3, column=col_idx)
            cell.value = header
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_alignment
            cell.border = border_style

        # Write data
        for row_idx, row_data in enumerate(df.itertuples(index=False), 4):
            ws.cell(row=row_idx, column=1, value=row_data.id).alignment = cell_alignment
            ws.cell(row=row_idx, column=2, value=str(row_data.date)[:10]).alignment = (
                cell_alignment
            )
            ws.cell(row=row_idx, column=3, value=row_data.category).alignment = (
                cell_alignment
            )
            ws.cell(row=row_idx, column=4, value=row_data.location).alignment = (
                cell_alignment
            )
            ws.cell(row=row_idx, column=5, value=row_data.price).alignment = (
                number_alignment
            )
            ws.cell(row=row_idx, column=5).number_format = "R$ #,##0.00"
            ws.cell(row=row_idx, column=6, value=row_data.payment_method).alignment = (
                cell_alignment
            )
            ws.cell(row=row_idx, column=7, value=row_data.installments).alignment = (
                number_alignment
            )

            # Apply borders
            for col_idx in range(1, 8):
                ws.cell(row=row_idx, column=col_idx).border = border_style

        # Add total row
        total_row = len(df) + 4
        ws.cell(row=total_row, column=1, value="TOTAL").font = Font(bold=True)
        ws.cell(row=total_row, column=1).alignment = cell_alignment
        ws.cell(row=total_row, column=5, value=df["price"].sum()).font = Font(bold=True)
        ws.cell(row=total_row, column=5).alignment = number_alignment
        ws.cell(row=total_row, column=5).number_format = "R$ #,##0.00"
        ws.cell(row=total_row, column=5).fill = PatternFill(
            start_color="E7E6E6", end_color="E7E6E6", fill_type="solid"
        )

        # Apply borders to total row
        for col_idx in range(1, 8):
            ws.cell(row=total_row, column=col_idx).border = border_style

        # Adjust column widths
        ws.column_dimensions["A"].width = 8
        ws.column_dimensions["B"].width = 15
        ws.column_dimensions["C"].width = 20
        ws.column_dimensions["D"].width = 25
        ws.column_dimensions["E"].width = 18
        ws.column_dimensions["F"].width = 18
        ws.column_dimensions["G"].width = 15

        # Save to buffer
        buffer = BytesIO()
        wb.save(buffer)
        buffer.seek(0)

        return buffer

    except Exception as e:
        raise Exception(f"Error creating formatted Excel: {e}")


def _show_analytics_view(expense_service: ExpenseService):
    """
    Show analytics and visualizations for expenses.
    DARK MODE PERMANENT.

    Args:
        expense_service: Service for expense operations
    """
    from frontend.components import (
        expense_pie_chart,
        expense_bar_chart,
        monthly_trend_chart,
        payment_method_chart,
        category_summary_table,
        metric_card,
    )

    st.markdown(
        """
        <h2 style="font-size: 24px; font-weight: 600; color: #ffffff; margin-bottom: 24px;">
            📊 Expense Analytics
        </h2>
        """,
        unsafe_allow_html=True,
    )

    try:
        # Get data
        df_expenses = expense_service.get_all_expenses()

        if df_expenses.empty:
            st.markdown(
                """
                <div style="padding: 48px; background: #1a1a1a; border-radius: 8px; 
                            text-align: center; border: 2px dashed #404040;">
                    <div style="font-size: 48px; margin-bottom: 16px;">📊</div>
                    <h3 style="color: #e0e0e0; font-size: 18px; font-weight: 600; margin: 0;">
                        No expense data available
                    </h3>
                    <p style="color: #808080; font-size: 14px; margin: 8px 0 0 0;">
                        Add some expenses to see analytics and visualizations
                    </p>
                </div>
                """,
                unsafe_allow_html=True,
            )
            return

        # Quick Stats
        stats = expense_service.get_expense_stats()

        st.markdown("### 📈 Quick Stats")
        col1, col2, col3 = st.columns(3)

        with col1:
            metric_card(
                title="Total Spent", value=f"R$ {stats.total_spent:,.2f}", icon="💸"
            )

        with col2:
            metric_card(
                title="Average Expense",
                value=f"R$ {stats.average_expense:,.2f}",
                icon="📊",
            )

        with col3:
            metric_card(title="Top Category", value=stats.top_category, icon="🏷️")

        st.divider()

        # Charts
        st.markdown("### 📈 Expense Analysis")

        # Get category data
        df_category = expense_service.get_expenses_by_category()

        col1, col2 = st.columns(2)

        with col1:
            expense_pie_chart(df_category, title="Spending by Category")

        with col2:
            expense_bar_chart(df_category, title="Category Breakdown")

        st.divider()

        # Payment method analysis
        st.markdown("### 💳 Payment Method Distribution")

        payment_dict = expense_service.get_expenses_by_payment_method()

        if payment_dict:
            payment_method_chart(payment_dict, title="Spending by Payment Method")
        else:
            st.info("No payment method data available")

        st.divider()

        # Monthly trend
        st.markdown("### 📅 Monthly Spending Trend")

        df_monthly = expense_service.get_monthly_expenses()

        if not df_monthly.empty:
            monthly_trend_chart(
                df_monthly, title="Monthly Expenses", y_label="Total Spent (R$)"
            )
        else:
            st.info("Not enough data for monthly trends")

        st.divider()

        # Category summary
        category_summary_table(df_category)

        st.divider()

        # Export section
        st.markdown("### 📥 Export Data")

        # Export formatted XLSX
        try:
            buffer = _create_formatted_expenses_xlsx(df_expenses)
            st.download_button(
                label="📊 Download Formatted Report (XLSX)",
                data=buffer,
                file_name=f"expenses_report_{datetime.now().strftime('%Y%m%d')}.xlsx",
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                use_container_width=True,
            )
        except Exception as e:
            st.error(f"❌ Error creating XLSX: {e}")

        st.divider()

        # Delete All section
        st.markdown("### 🗑️ Delete All Data")

        # Delete all button with confirmation
        if "show_delete_dialog_expenses" not in st.session_state:
            st.session_state.show_delete_dialog_expenses = False

        if not st.session_state.show_delete_dialog_expenses:
            if st.button(
                "🗑️ Delete All Expenses", type="secondary", use_container_width=True
            ):
                st.session_state.show_delete_dialog_expenses = True
                st.rerun()

        # Confirmation dialog
        if st.session_state.show_delete_dialog_expenses:
            st.warning(
                "⚠️ **Are you sure you want to delete ALL expenses?** This action cannot be undone!"
            )

            col1, col2 = st.columns(2)
            with col1:
                if st.button(
                    "❌ Cancel",
                    key="cancel_delete_all_expenses",
                    use_container_width=True,
                ):
                    st.session_state.show_delete_dialog_expenses = False
                    st.rerun()
            with col2:
                if st.button(
                    "✅ Confirm Delete",
                    key="confirm_delete_all_expenses",
                    type="primary",
                    use_container_width=True,
                ):
                    try:
                        expense_service.delete_all_expenses()
                        st.success("✅ All expenses deleted successfully!")
                        st.session_state.show_delete_dialog_expenses = False
                        st.rerun()
                    except Exception as e:
                        st.error(f"❌ Error deleting expenses: {e}")

    except Exception as e:
        st.error(f"❌ Error loading analytics: {e}")
